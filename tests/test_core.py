import openpyxl

from leadscraper import categories, osm
from leadscraper.enrich import extract_emails, extract_phones, normalize_url, _contact_links
from leadscraper.export import write_xlsx
from leadscraper.models import Lead
from leadscraper.pipeline import finalize, merge
from leadscraper.validate import best_phone, clean_email, rank_emails, verify_emails


def test_clean_email():
    assert clean_email("Info@Example-Shop.COM.") == "info@example-shop.com"
    assert clean_email("logo@2x.png") is None
    assert clean_email("noreply@shop.com") is None
    assert clean_email("a@example.com") is None
    assert clean_email("not an email") is None


def test_extract_emails_variants():
    html = ('<a href="mailto:sales@acme.pk?subject=Hi">x</a> contact: hello [at] acme [dot] pk '
            '<span data-cfemail="' + _cf("owner@acme.pk") + '"></span> <img src="a@2x.png">')
    assert set(extract_emails(html)) == {"sales@acme.pk", "hello@acme.pk", "owner@acme.pk"}


def _cf(email, key=0x42):
    return f"{key:02x}" + "".join(f"{ord(c) ^ key:02x}" for c in email)


def test_phone_normalisation():
    assert best_phone(["0321 4567890"], "PK")[0] == "+923214567890"
    assert best_phone(["12345"], "PK") == ("", "", "")
    assert best_phone(["+44 20 7946 0958; 000"], "PK")[0] == "+442079460958"


def test_extract_phones():
    html = '<a href="tel:+44 20 7946 0958">call</a> or +1 650-253-0000'
    ph = extract_phones(html, "GB")
    assert best_phone(ph, "GB")[0] == "+442079460958" and "+16502530000" in ph


def test_rank_prefers_same_domain_and_named():
    r = rank_emails(["info@gmail.com", "info@acme.com", "john@acme.com"], "https://www.acme.com")
    assert r[0] == "john@acme.com"


def test_verify_without_dns():
    best, others, status = verify_emails(["bad", "INFO@acme.com"], "acme.com", check_dns=False)
    assert best == "info@acme.com" and others == [] and "matches website" in status


def test_urls():
    assert normalize_url("acme.com") == "http://acme.com"
    assert normalize_url("http://localhost/x") == ""
    assert normalize_url("http://192.168.0.1") == ""


def test_contact_links():
    html = '<a href="/about-us">About</a><a href="/blog">Blog</a><a href="/contact">Contact</a>'
    links = _contact_links(html, "https://acme.com/", 3)
    assert links[0].endswith("/contact") and not any("blog" in l for l in links)


def test_categories():
    assert categories.resolve("Dentists")[1] is True
    assert categories.resolve("shop=bakery")[0] == ['["shop"="bakery"]']
    sel, exact = categories.resolve('weird "thing"')
    assert exact is False and '\\"' in sel[0]


def test_query_and_parse():
    place = osm.Place("Lahore", "PK", "Pakistan", "relation", 123, (0, 0, 1, 1))
    q, exact = osm.build_query(place, "dentist")
    assert "area(id:3600000123)" in q and exact and 'nwr["amenity"="dentist"](area.a);' in q
    els = [{"type": "node", "id": 1, "lat": 1.0, "lon": 2.0, "tags": {
        "name": "Smile Clinic", "phone": "+92 42 111 222 333", "website": "smile.pk",
        "email": "a@smile.pk;b@smile.pk", "addr:street": "Mall Rd", "addr:housenumber": "5",
        "addr:city": "Lahore"}}, {"type": "node", "id": 2, "tags": {"amenity": "dentist"}}]
    leads = osm.parse_elements(els, "dentist", "Lahore", "Pakistan")
    assert len(leads) == 1 and leads[0].address == "5 Mall Rd, Lahore" and leads[0].raw_emails == ["a@smile.pk", "b@smile.pk"]


def test_merge_dedupes():
    a = Lead("Smile Clinic", website="https://www.smile.pk", raw_phones=["1"])
    b = Lead("Smile Clinic", website="smile.pk", address="Mall Rd", source="Google Places", raw_phones=["2"])
    out = merge([a, b])
    assert len(out) == 1 and out[0].address == "Mall Rd" and out[0].raw_phones == ["1", "2"]


def test_export_roundtrip(tmp_path):
    l = Lead("=Evil", website="http://x.pk", raw_phones=["+442079460958"], raw_emails=["info@x.pk"])
    finalize([l], "GB", check_dns=False)
    p = tmp_path / "o.xlsx"
    write_xlsx([l], str(p), {"City": "X"})
    ws = openpyxl.load_workbook(p)["Leads"]
    assert ws["A2"].data_type == "s" and ws["C2"].value == "info@x.pk" and ws.max_column == 6 and ws["E2"].value == "Yes"


def test_crawl_with_fake_session():
    from leadscraper.enrich import crawl

    pages = {
        "http://shop.example.net": '<a href="/contact-us">Contact</a><a href="https://facebook.com/shop">f</a>',
        "http://shop.example.net/contact-us": 'Mail us: <a href="mailto:hi@shop.example.net">x</a> Tel +44 20 7946 0958',
    }

    class Raw:
        def __init__(self, b): self.b = b.encode()
        def read(self, n, decode_content=True): return self.b

    class Resp:
        def __init__(self, url):
            self.url, self.ok, self.encoding = url, url in pages, "utf-8"
            self.headers = {"content-type": "text/html"}
            self.raw = Raw(pages.get(url, ""))
        def close(self): pass

    class Sess:
        def get(self, url, **kw): return Resp(url)

    info = crawl("shop.example.net", "GB", Sess(), None)
    assert info.emails == ["hi@shop.example.net"] and info.socials["facebook"].endswith("/shop")
    assert any("7946" in p for p in info.phones)


def test_cli_end_to_end(tmp_path, monkeypatch):
    from leadscraper import cli

    def fake_search(city, country, cat, log=print):
        return [Lead("Smile Clinic", cat, "5 Mall Rd", city, country, raw_phones=["+92 42 35761234"],
                     raw_emails=["info@smile.pk"], lat=1, lon=1, source="OpenStreetMap"),
                Lead("No Contact Dental", cat, city=city, country=country, source="OpenStreetMap")], \
               osm.Place("x", "PK", country, "relation", 1, (0, 0, 1, 1))
    monkeypatch.setattr(cli.osm, "search", fake_search)
    out = tmp_path / "out.xlsx"
    rc = cli.main(["-c", "Pakistan", "-t", "Lahore", "-k", "dentist", "-o", str(out),
                   "--source", "osm", "--no-website-crawl", "--no-dns-check"])
    assert rc == 0
    ws = openpyxl.load_workbook(out)["Leads"]
    assert ws["A2"].value == "Smile Clinic" and ws["B2"].value == "+92 42 35761234" and ws["C2"].value == "info@smile.pk" and ws["C3"].value is None


def test_social_sites_not_crawled_and_region():
    from leadscraper.validate import country_region
    assert normalize_url("https://www.facebook.com/abc") == ""
    assert country_region("Pakistan") == "PK" and country_region("de") == "DE"
    assert country_region("Narnia") is None
    for name, code in {"UK": "GB", "UAE": "AE", "USA": "US", "Turkey": "TR", "South Korea": "KR",
                       "Germany": "DE", "India": "IN", "Brazil": "BR", "Nigeria": "NG"}.items():
        assert country_region(name) == code, name
    assert country_region("zz") is None


def test_maps_extract_js_markup():
    import glob, asyncio
    exe = glob.glob("/opt/pw-browsers/chromium-*/chrome-linux*/chrome")
    try:
        from playwright.async_api import async_playwright
    except ImportError:
        return
    if not exe:
        return
    from leadscraper.maps import EXTRACT_JS
    html = ('<h1>Results</h1><h1>Smile Hub</h1><button data-item-id="address" aria-label="Address: 5 Mall Rd">'
            '</button><button data-item-id="phone:tel:0320" aria-label="Phone: 0320 4411688"></button>'
            '<a data-item-id="authority" href="https://smilehub.pk/"></a>')

    async def go():
        async with async_playwright() as pw:
            b = await pw.chromium.launch(executable_path=exe[0], args=["--no-sandbox"])
            pg = await b.new_page()
            await pg.set_content(html)
            r = await pg.evaluate(EXTRACT_JS)
            await b.close()
            return r
    r = asyncio.run(go())
    assert r["name"] == "Smile Hub" and r["address"] == "5 Mall Rd" and r["phone"] == "0320 4411688"


# ---------- exact count / website status ----------

def _fake_osm(n, with_contact=lambda i: True):
    def fake_search(city, country, cat, log=print):
        leads = [Lead(f"Biz {i}", cat, f"{i} Main St", city, country, lat=i, lon=i, source="OpenStreetMap",
                      website=("https://facebook.com/biz%d" % i if i % 3 == 0 else "https://biz%d.com" % i if i % 3 == 1 else ""),
                      raw_phones=(["+44 20 7946 %04d" % i] if with_contact(i) else []))
                 for i in range(1, n + 1)]
        return leads, osm.Place("x", "GB", country, "relation", 1, (0, 0, 1, 1))
    return fake_search


def _run_cli(tmp_path, monkeypatch, fake, *extra):
    from leadscraper import cli
    monkeypatch.setattr(cli.osm, "search", fake)
    out = tmp_path / "o.xlsx"
    out.unlink(missing_ok=True)   # so a run that writes nothing is not mistaken for the previous file
    rc = cli.main(["-c", "UK", "-t", "London", "-k", "x", "-o", str(out), "--source", "osm",
                   "--no-website-crawl", "--no-dns-check", *extra])
    rows = list(openpyxl.load_workbook(out)["Leads"].iter_rows(min_row=2, values_only=True)) if out.exists() else []
    return rc, rows


def test_exact_count_even_when_half_have_no_contact(tmp_path, monkeypatch):
    # only even-numbered businesses have a phone; asking for 7 must give exactly 7 qualified leads
    rc, rows = _run_cli(tmp_path, monkeypatch, _fake_osm(100, lambda i: i % 2 == 0), "-n", "7")
    assert rc == 0 and len(rows) == 7 and all(r[1] for r in rows)


def test_count_larger_than_available_returns_what_exists(tmp_path, monkeypatch):
    rc, rows = _run_cli(tmp_path, monkeypatch, _fake_osm(5), "-n", "50")
    assert rc == 0 and len(rows) == 5


def test_website_status_column_and_filters(tmp_path, monkeypatch):
    rc, rows = _run_cli(tmp_path, monkeypatch, _fake_osm(9))
    status = {r[0]: r[4] for r in rows}
    link = {r[0]: r[5] for r in rows}
    assert status["Biz 1"] == "Yes" and link["Biz 1"] == "https://biz1.com"
    assert status["Biz 3"] == "No" and link["Biz 3"] is None      # Facebook page is not a website
    assert status["Biz 2"] == "No"
    _, rows = _run_cli(tmp_path, monkeypatch, _fake_osm(9), "--website", "no", "--include-seen")
    assert {r[4] for r in rows} == {"No"}
    _, rows = _run_cli(tmp_path, monkeypatch, _fake_osm(9), "--website", "yes", "--include-seen")
    assert {r[4] for r in rows} == {"Yes"}


def test_crawl_reports_reachable():
    from leadscraper.enrich import crawl, is_social

    class Resp:
        def __init__(self, ok): self.ok, self.url, self.encoding, self.headers = ok, "http://a.com", "utf-8", {"content-type": "text/html"}; self.raw = type("R", (), {"read": lambda s, n, decode_content=True: b"<p>hi</p>"})()
        def close(self): pass

    class Up:
        def get(self, url, **kw): return Resp(True)

    class Down:
        def get(self, url, **kw):
            import requests
            raise requests.ConnectionError()

    assert crawl("a.com", "GB", Up(), None).reachable is True
    assert crawl("a.com", "GB", Down(), None).reachable is False
    assert is_social("https://www.facebook.com/x") and not is_social("https://mysite.com")


# ---------- Google Maps streaming against a local mock of the Maps page ----------

def test_maps_scrolls_and_stops_at_exact_count(tmp_path, monkeypatch):
    import glob, threading
    from http.server import BaseHTTPRequestHandler, HTTPServer
    exe = glob.glob("/opt/pw-browsers/chromium-*/chrome-linux*/chrome")
    try:
        import playwright  # noqa: F401
    except ImportError:
        return
    if not exe:
        return
    monkeypatch.setenv("LEADSCRAPER_BROWSER_PATH", exe[0])
    TOTAL = 45
    hits = []

    class H(BaseHTTPRequestHandler):
        def log_message(self, *a): pass
        def do_GET(self):
            if self.path.startswith("/maps/search/"):
                body = """<html><body><h1>Results</h1>
<div role="feed" id="f" style="height:300px;overflow:auto">%s</div>
<script>
let n = 20; const T = %d, f = document.getElementById('f');
function add(a,b){ for(let i=a;i<b;i++){ const el=document.createElement('a'); el.href='/maps/place/Biz'+i+'/data=!1sID'+i; el.setAttribute('aria-label','Biz '+i); el.style.display='block'; el.style.height='60px'; el.textContent='Biz '+i; f.appendChild(el);} }
f.addEventListener('scroll', () => { if (f.scrollTop + f.clientHeight >= f.scrollHeight - 5 && n < T) { setTimeout(() => { const m=Math.min(T,n+20); add(n,m); n=m; if(n>=T){const d=document.createElement('div'); d.textContent="You've reached the end of the list."; f.appendChild(d);} }, 200); } });
</script></body></html>""" % ("".join('<a href="/maps/place/Biz%d/data=!1sID%d" aria-label="Biz %d" style="display:block;height:60px">Biz %d</a>' % (i, i, i, i) for i in range(20)), TOTAL)
            elif self.path.startswith("/maps/place/"):
                i = int(self.path.split("/")[3][3:])
                hits.append(i)
                phone = ('<button data-item-id="phone:tel:x" aria-label="Phone: 020 7946 %04d"></button>' % i) if i % 2 == 0 else ""
                site = '<a data-item-id="authority" href="https://biz%d.com/"></a>' % i if i % 4 == 0 else ""
                body = '<html><body><h1>Biz %d</h1><button data-item-id="address" aria-label="Address: %d High St"></button>%s%s</body></html>' % (i, i, phone, site)
            else:
                self.send_response(404); self.end_headers(); return
            self.send_response(200); self.send_header("Content-Type", "text/html"); self.end_headers()
            self.wfile.write(body.encode())

    srv = HTTPServer(("127.0.0.1", 0), H)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    from leadscraper import cli, maps
    monkeypatch.setattr(maps, "SEARCH_URL", "http://127.0.0.1:%d/maps/search/{q}" % srv.server_port)
    out = tmp_path / "m.xlsx"
    try:
        rc = cli.main(["-c", "UK", "-t", "London", "-k", "dentist", "-n", "15", "-o", str(out),
                       "--no-website-crawl", "--no-dns-check"])
        first = list(openpyxl.load_workbook(out)["Leads"].iter_rows(min_row=2, values_only=True))
        reads_before = len(hits)
        out2 = tmp_path / "m2.xlsx"
        rc2 = cli.main(["-c", "UK", "-t", "London", "-k", "dentist", "-n", "5", "-o", str(out2),
                        "--no-website-crawl", "--no-dns-check"])
        second = list(openpyxl.load_workbook(out2)["Leads"].iter_rows(min_row=2, values_only=True))
        second_reads = hits[reads_before:]
    finally:
        srv.shutdown()
    rows = first
    # second run: 5 brand-new leads, none repeated, and exported places were not even opened again
    assert rc2 == 0 and len(second) == 5 and not ({r[0] for r in rows} & {r[0] for r in second})
    exported_ids = {int(r[0].split()[1]) for r in rows}
    assert not (exported_ids & set(second_reads))
    assert rc == 0 and len(rows) == 15 and all(r[1] for r in rows)
    assert len(set(hits)) < TOTAL          # stopped early instead of reading every place
    assert max(hits) >= 25                 # had to scroll past the first 20 to find 15 with phones


def test_clean_website_drops_social_and_adds_scheme():
    from leadscraper.pipeline import clean_website
    for url, want in [("https://www.facebook.com/shop", ""), ("instagram.com/shop", ""),
                      ("smile.pk", "http://smile.pk"), ("https://smile.pk/", "https://smile.pk/"), ("", "")]:
        l = Lead("x", website=url)
        clean_website(l)
        assert l.website == want, url


# ---------- history: no repeats across runs ----------

def _two_runs(tmp_path, monkeypatch, n, *extra):
    r1 = _run_cli(tmp_path, monkeypatch, _fake_osm(30), "-n", str(n))[1]
    r2 = _run_cli(tmp_path, monkeypatch, _fake_osm(30), "-n", str(n), *extra)[1]
    return {r[0] for r in r1}, {r[0] for r in r2}


def test_second_run_returns_only_new_leads(tmp_path, monkeypatch):
    a, b = _two_runs(tmp_path, monkeypatch, 10)
    assert len(a) == 10 and len(b) == 10 and not (a & b)


def test_runs_until_source_exhausted_then_reports_no_new(tmp_path, monkeypatch):
    _run_cli(tmp_path, monkeypatch, _fake_osm(12), "-n", "12")
    rc, rows = _run_cli(tmp_path, monkeypatch, _fake_osm(12), "-n", "12")
    assert rc == 1 and rows == []


def test_include_seen_and_reset(tmp_path, monkeypatch):
    a, b = _two_runs(tmp_path, monkeypatch, 5, "--include-seen")
    assert a == b
    _run_cli(tmp_path, monkeypatch, _fake_osm(30), "-n", "5")
    _, c = _run_cli(tmp_path, monkeypatch, _fake_osm(30), "-n", "5", "--reset-history")
    assert {r[0] for r in c} == a


def test_same_phone_or_email_counts_as_same_lead(tmp_path, monkeypatch):
    def fake(city, country, cat, log=print):
        leads = [Lead("Clinic A", cat, lat=1, lon=1, raw_phones=["+44 20 7946 0001"]),
                 Lead("Clinic A (listing 2)", cat, lat=9, lon=9, raw_phones=["+44 20 7946 0001"]),
                 Lead("Other", cat, lat=5, lon=5, raw_emails=["hi@other.co.uk"])]
        return leads, osm.Place("x", "GB", country, "relation", 1, (0, 0, 1, 1))
    _, rows = _run_cli(tmp_path, monkeypatch, fake)
    assert len(rows) == 2


def test_history_file_survives_corruption(tmp_path):
    from leadscraper.history import History
    p = tmp_path / "h.json"
    p.write_text("{not json")
    h = History(p)
    assert h.exported == 0 and (tmp_path / "h.broken").exists()


# ---------- email reasons + recall fixes ----------

def test_mailto_url_encoding_fixed():
    assert clean_email("%20info@smileon.pk") == "info@smileon.pk"
    assert extract_emails('<a href="mailto:%20info@smileon.pk">x</a>') == ["info@smileon.pk"]
    assert set(extract_emails('<a href="mailto:a@x.pk,b@x.pk?subject=Hi">x</a>')) == {"a@x.pk", "b@x.pk"}


def test_embedded_json_ld_email_only_if_same_domain():
    from leadscraper.enrich import extract_embedded_emails, crawl
    html = ('<script type="application/ld+json">{"@type":"Dentist","email":"care@smile.pk"}</script>'
            '<script>var s="support@elementor.com"; var u="owner\\u0040smile.pk";</script>')
    assert set(extract_embedded_emails(html)) == {"care@smile.pk", "support@elementor.com", "owner@smile.pk"}

    class Raw:
        def read(self, n, decode_content=True): return html.encode()

    class Resp:
        ok, url, encoding, headers, raw = True, "http://smile.pk", "utf-8", {"content-type": "text/html"}, Raw()
        def close(self): pass

    class S:
        def get(self, url, **kw): return Resp()

    # visible text has no email; the third-party address in the script must NOT be trusted
    assert set(crawl("smile.pk", "PK", S(), None).emails) == {"care@smile.pk", "owner@smile.pk"}


def test_email_reasons(tmp_path):
    cases = {
        "none": Lead("A"),
        "ok": Lead("B", website="http://b.pk", site_state="ok"),
        "down": Lead("C", website="http://c.pk", site_state="down"),
        "blocked": Lead("D", website="http://d.pk", site_state="blocked"),
        "unchecked": Lead("E", website="http://e.pk"),
        "bad": Lead("F", website="http://f.pk", site_state="ok", raw_emails=["x@no-such-domain.invalid"]),
        "good": Lead("G", website="http://g.pk", site_state="ok", raw_emails=["hi@g.pk"]),
    }
    leads = list(cases.values())
    finalize(leads, "PK", check_dns=False)
    notes = {k: v.email_note for k, v in cases.items()}
    assert notes["none"] == "No website - email not available"
    assert notes["ok"] == "Not found on website"
    assert notes["down"] == "Website not responding"
    assert notes["blocked"] == "Website blocks automated access"
    assert notes["unchecked"] == "Not checked"
    assert cases["good"].email == "hi@g.pk" and cases["good"].email_note == ""
    p = tmp_path / "n.xlsx"
    write_xlsx(leads, str(p), {})
    ws = openpyxl.load_workbook(p)["Leads"]
    col = {ws.cell(r, 1).value: ws.cell(r, 3).value for r in range(2, ws.max_row + 1)}
    assert col["A"] == "No website - email not available" and col["G"] == "hi@g.pk"
    assert ws.cell(2, 3).hyperlink is None and ws.cell(8, 3).hyperlink is not None
