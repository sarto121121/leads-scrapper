from __future__ import annotations

import re
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Lead

COLUMNS = [("Name", 38), ("Phone", 22), ("Email", 38), ("Address", 60), ("Website Present", 18), ("Website", 40)]


def _row(l: Lead) -> list:
    return [l.name, l.phone_display, l.email, l.address, l.website_status or ('Yes' if l.website else 'No'), l.website]


def safe_filename(*parts: str) -> str:
    return re.sub(r"[^\w\-]+", "_", "_".join(parts)).strip("_").lower()[:80]


def write_xlsx(leads: list[Lead], path: str, meta: dict[str, str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Leads"
    ws.append([c for c, _ in COLUMNS])
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(vertical="center")
    for i, (_, w) in enumerate(COLUMNS, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for l in leads:
        ws.append([_clip(v) for v in _row(l)])
    for r in range(2, ws.max_row + 1):
        c = ws.cell(r, 3)
        if c.value:
            c.hyperlink = "mailto:" + str(c.value)
            c.font = Font(color="0563C1", underline="single")
        w = ws.cell(r, 6)
        if w.value and str(w.value).startswith("http"):
            w.hyperlink = str(w.value)
            w.font = Font(color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions

    s = wb.create_sheet("Summary")
    n = len(leads)
    stats = [
        ("Generated", datetime.now().strftime("%Y-%m-%d %H:%M")), *meta.items(),
        ("Total leads", n),
        ("With email", sum(bool(l.email) for l in leads)),
        ("With phone", sum(bool(l.phone) for l in leads)),
        ("With a website", sum(l.website_status.startswith("Yes") for l in leads)),
        ("Social page only", sum(l.website_status == "Social page only" for l in leads)),
        ("No website", sum(l.website_status == "No" for l in leads)),
        ("With email AND phone", sum(bool(l.email and l.phone) for l in leads)),
        ("Note", "Emails are only listed when found on the business's own listing/website and the "
                 "domain accepts mail. Phones are only listed when they are valid for their country."),
    ]
    for k, v in stats:
        s.append([k, v])
        s.cell(s.max_row, 1).font = Font(bold=True)
    s.column_dimensions["A"].width = 24
    s.column_dimensions["B"].width = 100
    wb.save(path)


def _clip(v):
    """Excel cells hold at most 32767 chars; also neutralise formula injection from scraped text."""
    if isinstance(v, str):
        v = v[:32000]
        if v[:1] in ("=", "+", "-", "@") and not v.startswith("+"):
            v = " " + v
    return v
